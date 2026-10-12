# ============================================================
# PANEL ADMIN — ÉTAPE 11 : SALAIRE DÉSAGRÉGÉ PAR ANNÉES D'EXPÉRIENCE
# ============================================================
# Désagrège le salaire par tranches d'années d'expérience
# (classes de 5 ans) calculées depuis la date de prise de service.
#
#   annees_experience = (période d'observation − prise_service) / 365,25
#
# Ventilations produites : global, par année, par sexe, par grade,
# par statut, par organisme/secteur, par emploi, par fonction,
# par service, par lieu d'affectation, par poste, par CITP,
# par situation matrimoniale.
#
# Source  : s3://gold/panel_admin/panel_complet.parquet
# Sortie  : s3://staging/panel_admin/exports_gold/salaire_experience.xlsx
#
# Structure du classeur produit :
#   SOMMAIRE        — définitions M1/M2 + navigation cliquable
#   Métadonnées     — méthode, périmètre, exclusions
#   18 feuilles     — une par ventilation (bloc M1 puis bloc M2)
#                     classes de 5 ans : [0-5[, [5-10[, … [40-45[, [45+
#
# VARIABLES SALAIRE (cf. 07b_calcul_indicateurs.py)
# -------------------------------------------------
#   M1 = montant_brut_sans_ind_sans_alloc
#        Population « Statut conforme (barème) » ; codes postes
#        124+228+271+425 ; HORS indemnités et HORS allocation familiale (421).
#   M2 = montant_brut_total
#        Population « Statut_Total conforme (barème) » ; Feuille 1,
#        tous codes postes confondus.
#
# Usage :
#   python 11_analyse_salaire_experience.py            # réutilise le cache local si présent
#   python 11_analyse_salaire_experience.py --force    # reconstruit le cache
#
# Dépendances :
#   pip install polars boto3 python-dotenv openpyxl pyarrow
# ============================================================

import argparse
import os
import tempfile
import time
from pathlib import Path

from dotenv import load_dotenv
import boto3
from botocore.client import Config
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

load_dotenv(".env")

# ============================================================
# CONFIGURATION
# ============================================================

MINIO_ENDPOINT   = os.getenv("MINIO_ENDPOINT")
MINIO_ACCESS_KEY = os.getenv("MINIO_ACCESS_KEY")
MINIO_SECRET_KEY = os.getenv("MINIO_SECRET_KEY")

BUCKET_GOLD    = os.getenv("BUCKET_GOLD", "gold")
BUCKET_STAGING = os.getenv("BUCKET_STAGING", "staging")

KEY_PANEL      = "panel_admin/panel_complet.parquet"
PREFIX_EXPORTS = "panel_admin/exports_gold"
OUTPUT_KEY     = f"{PREFIX_EXPORTS}/salaire_experience.xlsx"

ANNEES          = list(range(2015, 2027))
CLASSES_LARGEUR = [5]        # classes de 5 ans uniquement
CLASSE_PLAFOND  = 45         # au-delà : "45+"
DAYS_PER_YEAR   = 365.25

OUTDIR   = Path("analyse_experience_out")   # copie locale (inspection)
CACHE    = OUTDIR / "_cache_compact.parquet"
FICHIER_XLSX = OUTDIR / "salaire_experience.xlsx"

COLS_SRC = [
    "matricule", "annee", "mois", "prise_service", "sexe",
    "montant_brut_sans_ind_sans_alloc",     # M1
    "montant_brut_total",                   # M2
    # dimensions
    "CODE_ORGANISME", "organisme",
    "CODE_EMPLOI", "emploi",
    "CODE_GRADE", "GRADE",
    "statut_fonctionnaire",
    "CODE_FONCTION", "fonction",
    "CODE_SERVICE", "service",
    "CODE_AFFECTATION", "lieu_affectation",
    "CODE_POSTE",
    "Code_CITP", "Metier_CITP",
    "situation_matrimoniale",
]

METRIQUES = {
    "M1_sans_alloc": "m1",   # montant_brut_sans_ind_sans_alloc
    "M2_brut_total": "m2",   # montant_brut_total
}

LIB_METRIQUE = {
    "M1_sans_alloc": "M1 — montant_brut_sans_ind_sans_alloc  (hors indemnités, hors allocation familiale)",
    "M2_brut_total": "M2 — montant_brut_total  (Feuille 1, tous codes postes)",
}

# Ventilations : (nom de feuille, colonnes de groupe, description, section)
GROUPES = [
    ("Global",                   [],                          "Toutes populations confondues",                 "Synthèse"),
    ("ParAnnee",                 ["annee"],                   "Par année d'observation",                       "Évolution dans le temps"),
    ("ParSexe",                  ["sexe"],                    "Par sexe",                                      "Démographie"),
    ("ParSituationMatrimoniale", ["situation_matrimoniale"],  "Par situation matrimoniale",                    "Démographie"),
    ("ParGrade",                 ["GRADE"],                   "Par grade normalisé",                           "Grade & statut"),
    ("ParCodeGrade",             ["CODE_GRADE"],              "Par code grade",                                "Grade & statut"),
    ("ParStatut",                ["statut_fonctionnaire"],    "Par statut / catégorie de fonctionnaire",       "Grade & statut"),
    ("ParOrganisme",             ["organisme"],               "Par secteur / organisme (libellé)",             "Secteur / organisme"),
    ("ParCodeOrganisme",         ["code_organisme"],          "Par code organisme",                            "Secteur / organisme"),
    ("ParLieuAffectation",       ["lieu_affectation"],        "Par lieu d'affectation",                        "Secteur / organisme"),
    ("ParCodePoste",             ["CODE_POSTE"],              "Par code poste",                                "Secteur / organisme"),
    ("ParEmploi",                ["emploi"],                  "Par emploi",                                    "Emploi, fonction & service"),
    ("ParCodeEmploi",            ["CODE_EMPLOI"],             "Par code emploi",                               "Emploi, fonction & service"),
    ("ParFonction",              ["fonction"],                "Par fonction",                                  "Emploi, fonction & service"),
    ("ParCodeFonction",          ["CODE_FONCTION"],           "Par code fonction",                             "Emploi, fonction & service"),
    ("ParService",               ["service"],                 "Par service",                                   "Emploi, fonction & service"),
    ("ParCodeService",           ["CODE_SERVICE"],            "Par code service",                              "Emploi, fonction & service"),
    ("ParCITP",                  ["Code_CITP", "Metier_CITP"], "Par code et libellé CITP (métier)",             "Emploi, fonction & service"),
]

# ============================================================
# CLIENT S3
# ============================================================

s3 = boto3.client(
    "s3",
    endpoint_url          = MINIO_ENDPOINT,
    aws_access_key_id     = MINIO_ACCESS_KEY,
    aws_secret_access_key = MINIO_SECRET_KEY,
    config                = Config(
        signature_version = "s3v4",
        retries           = {"max_attempts": 5, "mode": "adaptive"},
        connect_timeout   = 30,
        read_timeout      = 300,
    ),
    region_name = "us-east-1",
    verify      = False,
)

# ============================================================
# OUTILS POLARS
# ============================================================

def parse_date_multi(col: str) -> pl.Expr:
    s = pl.col(col).cast(pl.Utf8, strict=False)
    return pl.coalesce([
        s.str.to_date("%Y-%m-%d", strict=False),
        s.str.to_date("%d/%m/%Y", strict=False),
        s.str.to_date("%d-%m-%Y", strict=False),
        s.str.to_date("%Y/%m/%d", strict=False),
        s.str.to_date("%d.%m.%Y", strict=False),
        s.str.to_datetime(strict=False).cast(pl.Date),
    ])


def plafond(largeur: int) -> int:
    return (CLASSE_PLAFOND // largeur) * largeur


def expr_classe(largeur: int) -> pl.Expr:
    haut = plafond(largeur)
    borne = ((pl.col("annees_experience") / largeur).floor() * largeur).clip(0, haut)
    return (
        pl.when(borne >= haut)
          .then(pl.lit(f"{haut}+"))
          .otherwise(pl.format("{} - {}", borne.cast(pl.Int32), (borne + largeur).cast(pl.Int32)))
    )


def expr_borne_inf(largeur: int) -> pl.Expr:
    return ((pl.col("annees_experience") / largeur).floor() * largeur).clip(0, plafond(largeur))

# ============================================================
# TÉLÉCHARGEMENT + CACHE LOCAL COMPACT
# ============================================================

def telecharger_panel() -> Path:
    """Télécharge panel_complet.parquet (Gold) dans un fichier temporaire local."""
    info = s3.head_object(Bucket=BUCKET_GOLD, Key=KEY_PANEL)
    taille = info["ContentLength"]
    tmp = Path(tempfile.gettempdir()) / "panel_complet.parquet"
    print(f"  Téléchargement de s3://{BUCKET_GOLD}/{KEY_PANEL}  ({taille/1024**2:.1f} MB)")
    t0 = time.time()
    s3.download_file(BUCKET_GOLD, KEY_PANEL, str(tmp))
    print(f"  ✓ Téléchargé en {time.time()-t0:.1f}s → {tmp}")
    return tmp


def construire_cache(chemin_parquet: Path, force: bool) -> None:
    if CACHE.exists() and not force:
        print(f"  ↺ Cache local réutilisé : {CACHE}  (--force pour reconstruire)")
        return

    OUTDIR.mkdir(parents=True, exist_ok=True)
    t0 = time.time()

    lf = (
        pl.scan_parquet(str(chemin_parquet))
        .filter(pl.col("annee").is_in(ANNEES))
        .with_columns([
            parse_date_multi("prise_service").alias("_ps_dt"),
            pl.date(pl.col("annee").cast(pl.Int32),
                    pl.col("mois").cast(pl.Int32), 1).alias("_periode"),
        ])
        .with_columns(
            ((pl.col("_periode") - pl.col("_ps_dt")).dt.total_days() / DAYS_PER_YEAR)
            .cast(pl.Float32).alias("annees_experience")
        )
        .select([
            pl.col("matricule").cast(pl.Utf8),
            pl.col("annee").cast(pl.Int32),
            pl.col("sexe").cast(pl.Utf8),
            pl.col("annees_experience"),
            pl.col("montant_brut_sans_ind_sans_alloc").cast(pl.Float64).alias("m1"),
            pl.col("montant_brut_total").cast(pl.Float64).alias("m2"),
            pl.col("CODE_ORGANISME").cast(pl.Utf8).alias("code_organisme"),
            pl.col("organisme").cast(pl.Utf8),
            pl.col("CODE_EMPLOI").cast(pl.Utf8),
            pl.col("emploi").cast(pl.Utf8),
            pl.col("CODE_GRADE").cast(pl.Utf8),
            pl.col("GRADE").cast(pl.Utf8),
            pl.col("statut_fonctionnaire").cast(pl.Utf8),
            pl.col("CODE_FONCTION").cast(pl.Utf8),
            pl.col("fonction").cast(pl.Utf8),
            pl.col("CODE_SERVICE").cast(pl.Utf8),
            pl.col("service").cast(pl.Utf8),
            pl.col("CODE_AFFECTATION").cast(pl.Utf8),
            pl.col("lieu_affectation").cast(pl.Utf8),
            pl.col("CODE_POSTE").cast(pl.Utf8),
            pl.col("Code_CITP").cast(pl.Utf8),
            pl.col("Metier_CITP").cast(pl.Utf8),
            pl.col("situation_matrimoniale").cast(pl.Utf8),
        ])
    )

    # sink_parquet : écriture en streaming (mémoire maîtrisée, footer correct)
    lf.sink_parquet(str(CACHE), compression="zstd")

    n = pl.scan_parquet(str(CACHE)).select(pl.len()).collect().item()
    print(f"  ✓ Cache écrit : {CACHE}  ({n:,} lignes, {time.time()-t0:.1f}s)")


def statistiques_qualite(chemin_parquet: Path) -> None:
    lf = pl.scan_parquet(str(chemin_parquet)).filter(pl.col("annee").is_in(ANNEES))
    lf = lf.with_columns([
        parse_date_multi("prise_service").alias("_ps_dt"),
        pl.date(pl.col("annee").cast(pl.Int32),
                pl.col("mois").cast(pl.Int32), 1).alias("_periode"),
    ])
    res = lf.select([
        pl.len().alias("n"),
        pl.col("_ps_dt").is_null().sum().alias("ps_null"),
        ((pl.col("_periode") - pl.col("_ps_dt")).dt.total_days() < 0).sum().alias("exp_neg"),
    ]).collect().row(0, named=True)
    print(f"    lignes (2015-2026) : {res['n']:,}")
    print(f"    prise_service manquante : {res['ps_null']:,}")
    print(f"    expérience négative (période < prise de service) : {res['exp_neg']:,}")

# ============================================================
# AGRÉGATIONS
# ============================================================

def tableau_stats(lf: pl.LazyFrame, largeur: int, metrique: str, groupe: list[str]) -> pl.DataFrame:
    col_val = METRIQUES[metrique]
    base = (
        lf
        .filter(pl.col("annees_experience").is_not_null() & (pl.col("annees_experience") >= 0))
        .filter(pl.col(col_val).is_not_null() & (pl.col(col_val) > 0))
        .with_columns(
            [expr_classe(largeur).alias("classe_experience"), expr_borne_inf(largeur).alias("_borne")]
            + [pl.col(c).cast(pl.Utf8).fill_null("(NC)").alias(c) for c in groupe]
        )
    )
    par = ["_borne", "classe_experience"] + groupe
    tre = ["_borne"] + groupe

    return (
        base
        .group_by(par)
        .agg([
            pl.len().alias("n_obs"),
            pl.col("matricule").n_unique().alias("n_agents"),
            pl.col(col_val).mean().round(0).alias("salaire_moyen"),
            pl.col(col_val).median().round(0).alias("salaire_median"),
            pl.col(col_val).quantile(0.25, interpolation="linear").round(0).alias("p25"),
            pl.col(col_val).quantile(0.75, interpolation="linear").round(0).alias("p75"),
            pl.col(col_val).min().round(0).alias("min"),
            pl.col(col_val).max().round(0).alias("max"),
        ])
        .sort(tre)
        .drop("_borne")
        .collect()
    )

# ============================================================
# ÉCRITURE EXCEL (SOMMAIRE + FEUILLES)
# ============================================================

ENTETES = {
    "classe_experience": "Classe d'expérience",
    "annee": "Année",
    "sexe": "Sexe",
    "situation_matrimoniale": "Situation matrimoniale",
    "GRADE": "Grade",
    "CODE_GRADE": "Code grade",
    "statut_fonctionnaire": "Statut",
    "organisme": "Secteur / organisme",
    "code_organisme": "Code organisme",
    "lieu_affectation": "Lieu d'affectation",
    "CODE_POSTE": "Code poste",
    "emploi": "Emploi",
    "CODE_EMPLOI": "Code emploi",
    "fonction": "Fonction",
    "CODE_FONCTION": "Code fonction",
    "service": "Service",
    "CODE_SERVICE": "Code service",
    "Code_CITP": "Code CITP",
    "Metier_CITP": "Métier (CITP)",
    "n_obs": "Observations",
    "n_agents": "Agents distincts",
    "salaire_moyen": "Salaire moyen",
    "salaire_median": "Salaire médian",
    "p25": "P25",
    "p75": "P75",
    "min": "Min",
    "max": "Max",
}

_FILL_TITRE  = PatternFill("solid", fgColor="1F4E78")
_FILL_ENTETE = PatternFill("solid", fgColor="2E75B6")
_FILL_SOUS   = PatternFill("solid", fgColor="DDEBF7")
_FONT_TITRE  = Font(color="FFFFFF", bold=True, size=12)
_FONT_ENTETE = Font(color="FFFFFF", bold=True)
_FONT_SOUS   = Font(bold=True, color="1F4E78")
_FONT_LIEN   = Font(color="0563C1", underline="single")
_ALIGN_C     = Alignment(horizontal="center", vertical="center")
_ALIGN_WRAP  = Alignment(wrap_text=True, vertical="top")


def _cell(ws, r, c, v, font=None, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    if font:  cell.font = font
    if fill:  cell.fill = fill
    if align: cell.alignment = align
    return cell


def _titre_feuille(ws, texte, n_cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n_cols, 4))
    _cell(ws, 1, 1, texte, font=_FONT_TITRE, fill=_FILL_TITRE, align=_ALIGN_C)
    ws.row_dimensions[1].height = 26


def _lien(ws, r, c, texte, cible):
    cible_safe = cible.replace("'", "''")
    _cell(ws, r, c, f'=HYPERLINK("#{cible_safe}!A1","{texte}")', font=_FONT_LIEN)


def _ecrire_bloc(ws, r, sous_titre, df) -> int:
    n_cols = len(df.columns)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(n_cols, 4))
    _cell(ws, r, 1, sous_titre, font=_FONT_SOUS, fill=_FILL_SOUS,
          align=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[r].height = 20
    r += 1
    for j, c in enumerate(df.columns, start=1):
        _cell(ws, r, j, ENTETES.get(c, c), font=_FONT_ENTETE, fill=_FILL_ENTETE, align=_ALIGN_C)
    r += 1
    for ligne in df.iter_rows(named=False):
        for j, v in enumerate(ligne, start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    return r


def _ajuster_largeurs(ws, n_cols):
    for j in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14


def ecrire_feuille_groupe(wb, nom_feuille, titre, description, blocs):
    ws = wb.create_sheet(nom_feuille)
    n_cols = max((len(df.columns) for _, df in blocs), default=4)
    _titre_feuille(ws, titre, n_cols)
    _lien(ws, 2, 1, "← Retour au sommaire", "SOMMAIRE")
    _cell(ws, 3, 1, description, font=Font(italic=True, color="595959"))
    r = 5
    for sous_titre, df in blocs:
        r = _ecrire_bloc(ws, r, sous_titre, df)
        r += 1
    _ajuster_largeurs(ws, n_cols)
    ws.freeze_panes = "A5"


def ecrire_sommaire(wb, sections):
    ws = wb.create_sheet("SOMMAIRE")
    ws.merge_cells("A1:F1")
    _cell(ws, 1, 1, "SOMMAIRE — Salaire par années d'expérience (classes de 5 ans)",
          font=Font(color="FFFFFF", bold=True, size=14), fill=_FILL_TITRE, align=_ALIGN_C)
    ws.row_dimensions[1].height = 30

    r = 3
    _cell(ws, r, 1, "DÉFINITIONS DES VARIABLES", font=_FONT_SOUS); r += 1
    for cle, lib in [
        ("M1", "montant_brut_sans_ind_sans_alloc — population « Statut conforme (barème) » ; codes postes 124+228+271+425 ; HORS indemnités et HORS allocation familiale (421)."),
        ("M2", "montant_brut_total — population « Statut_Total conforme (barème) » ; Feuille 1, tous codes postes confondus."),
        ("Expérience", "annees_experience = (période d'observation − prise_service) / 365,25 ; classes de 5 ans : [0-5[, [5-10[, … [40-45[, [45+."),
        ("Source", f"s3://{BUCKET_GOLD}/{KEY_PANEL} (2015→2026)."),
        ("Exclusions", "salaires ≤ 0 exclus par métrique ; expérience négative et prise_service manquante exclues."),
    ]:
        _cell(ws, r, 1, cle, font=Font(bold=True))
        _cell(ws, r, 2, lib, align=_ALIGN_WRAP)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    _cell(ws, r, 1, "NAVIGATION", font=_FONT_SOUS); r += 1
    for j, h in enumerate(["Section", "Feuille (cliquer)", "Description", "Lignes (M1)"], start=1):
        _cell(ws, r, j, h, font=_FONT_ENTETE, fill=_FILL_ENTETE, align=_ALIGN_C)
    r += 1
    for section, entrees in sections.items():
        _cell(ws, r, 1, section, font=Font(bold=True, color="1F4E78"))
        r += 1
        for nom_feuille, titre, desc, n_lignes in entrees:
            _lien(ws, r, 2, nom_feuille, nom_feuille)
            _cell(ws, r, 3, desc)
            _cell(ws, r, 4, n_lignes, align=_ALIGN_C)
            r += 1
        r += 1
    _cell(ws, r, 1, "Métadonnées", font=Font(bold=True, color="1F4E78")); r += 1
    _lien(ws, r, 2, "Métadonnées", "Métadonnées")
    _cell(ws, r, 3, "Méthode, périmètre, exclusions, unités")
    r += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 14
    ws.freeze_panes = "A3"


def ecrire_metadonnees(wb):
    ws = wb.create_sheet("Métadonnées")
    _titre_feuille(ws, "MÉTADONNÉES", 4)
    _lien(ws, 2, 1, "← Retour au sommaire", "SOMMAIRE")
    lignes = [
        ("M1", "Méthodo 1 — Population « Statut conforme (barème) ». Variable : montant_brut_sans_ind_sans_alloc (codes postes 124+228+271+425, hors allocation familiale 421)."),
        ("M2", "Méthodo 2 — Population « Statut_Total conforme (barème) ». Variable : montant_brut_total (Feuille 1, tous codes postes confondus)."),
        ("Source", f"s3://{BUCKET_GOLD}/{KEY_PANEL} — lecture ; sortie dans s3://{BUCKET_STAGING}/{OUTPUT_KEY}."),
        ("Calcul expérience", "annees_experience = (période d'observation − prise_service) / 365,25 ; arrondi à la classe de 5 ans inférieure."),
        ("Classes (5 ans)", "[0-5[, [5-10[, [10-15[, [15-20[, [20-25[, [25-30[, [30-35[, [35-40[, [40-45[, [45+."),
        ("Exclusions M1", "montant_brut_sans_ind_sans_alloc ≤ 0 ou null."),
        ("Exclusions M2", "montant_brut_total ≤ 0 ou null."),
        ("Expérience négative", "période < prise de service → exclues du calcul."),
        ("prise_service manquante", "exclue du calcul d'expérience."),
        ("Statistiques", "n_obs = observations (agent-mois) ; n_agents = matricules distincts du groupe ; salaire moyen/médian, P25, P75, min, max (FCFA)."),
        ("Note secteur", "« Secteur / organisme » = champ organisme du panel (libellé). Le SECTEUR ANSTAT (Barême Général, ENS SUP, …) provient de la grille salariale et nécessite une jointure référentielle."),
    ]
    r = 4
    for cle, val in lignes:
        _cell(ws, r, 1, cle, font=Font(bold=True))
        _cell(ws, r, 2, val, align=_ALIGN_WRAP)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
        r += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100


def exporter(blocs_par_feuille: dict, ordre_sommaire: list) -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    sections: dict[str, list] = {}
    for section, nom, titre, desc, n_lignes in ordre_sommaire:
        sections.setdefault(section, []).append((nom, titre, desc, n_lignes))
    ecrire_sommaire(wb, sections)          # SOMMAIRE créé en premier

    for nom, (titre, desc, blocs) in blocs_par_feuille.items():
        ecrire_feuille_groupe(wb, nom, titre, desc, blocs)

    ecrire_metadonnees(wb)

    wb.save(FICHIER_XLSX)
    print(f"  ✓ Excel local : {FICHIER_XLSX}")

    # Upload vers MinIO
    s3.upload_file(
        str(FICHIER_XLSX), BUCKET_STAGING, OUTPUT_KEY,
        ExtraArgs={"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )
    print(f"  ✓ Uploadé : s3://{BUCKET_STAGING}/{OUTPUT_KEY}")

# ============================================================
# MAIN
# ============================================================

def main() -> None:
    ap = argparse.ArgumentParser(description="Salaire désagrégé par années d'expérience.")
    ap.add_argument("--force", action="store_true", help="Reconstruire le cache local")
    ap.add_argument("--garder-parquet", action="store_true",
                    help="Ne pas supprimer le panel_complet.parquet téléchargé")
    args = ap.parse_args()

    print("=" * 70)
    print(" ÉTAPE 11 — SALAIRE PAR ANNÉES D'EXPÉRIENCE (classes de 5 ans)")
    print(f" Source : s3://{BUCKET_GOLD}/{KEY_PANEL}")
    print(f" Sortie : s3://{BUCKET_STAGING}/{OUTPUT_KEY}")
    print("=" * 70)

    if CACHE.exists() and not args.force:
        print(f"  ↺ Cache local réutilisé : {CACHE}  (--force pour reconstruire)")
    else:
        chemin = telecharger_panel()
        try:
            statistiques_qualite(chemin)
            construire_cache(chemin, force=True)
        finally:
            if not args.garder_parquet and chemin.exists():
                chemin.unlink()

    lf = pl.scan_parquet(str(CACHE))

    blocs_par_feuille: dict[str, tuple] = {}
    ordre_sommaire: list[tuple] = []
    largeur = CLASSES_LARGEUR[0]

    for nom_feuille, groupe, desc, section in GROUPES:
        blocs = []
        for metrique in METRIQUES:
            df = tableau_stats(lf, largeur, metrique, groupe)
            blocs.append((LIB_METRIQUE[metrique], df))
        blocs_par_feuille[nom_feuille] = (nom_feuille, desc, blocs)
        ordre_sommaire.append((section, nom_feuille, nom_feuille, desc, len(blocs[0][1])))
        print(f"  · {nom_feuille} : {len(blocs[0][1])} lignes (M1)")

    exporter(blocs_par_feuille, ordre_sommaire)

    print("\n" + "=" * 70)
    print(" APERÇU — Global · M1 sans allocation familiale")
    print("=" * 70)
    with pl.Config(tbl_rows=15, tbl_cols=12, fmt_str_lengths=22):
        print(blocs_par_feuille["Global"][2][0][1])

    print("\n" + "=" * 70)
    print(f" ✅ TERMINÉ — {len(blocs_par_feuille)} ventilations × 2 métriques")
    print(f"    s3://{BUCKET_STAGING}/{OUTPUT_KEY}")
    print("=" * 70 + "\n")


if __name__ == "__main__":
    main()
