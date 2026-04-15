# ============================================================
# PANEL ADMIN — ÉTAPE 7c : ANALYSE DE LA COHÉRENCE DES GRADES
# ============================================================
# Lit le panel Gold complet et analyse la progression des grades
# par agent dans le temps.
#
# Hiérarchie des grades (du plus bas au plus haut) :
#   D1 → D2 → D3 → C1 → C2 → C3 → B1 → B3 → A3 → A4 → A5 → A6 → A7
#   (pas de B2 ; A est la catégorie supérieure)
#
# Feuilles produites :
#   RESUME_GLOBAL      : Stats par type de transition × source du grade
#   TRANSITIONS_DETAIL : Toutes les paires (grade_de × grade_à) × source × effectif
#   AMPLITUDE_SAUTS    : Distribution des amplitudes de saut (en niveaux)
#
# Prérequis : 06_compiler_panel.py doit avoir été exécuté.
#
# Source  : s3://gold/panel_admin/panel_YYYY.parquet
# Sortie  : s3://staging/panel_admin/exports_gold/coherence_grades.xlsx
# ============================================================

import io
import os
import re
from dotenv import load_dotenv
import boto3
from botocore.client import Config
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

load_dotenv(".env")

# ============================================================
# CONNEXION S3 / MINIO
# ============================================================

MINIO_ENDPOINT   = os.getenv("MINIO_ENDPOINT",   "http://192.168.1.230:30137")
MINIO_ACCESS_KEY = os.getenv("MINIO_ACCESS_KEY", "datalab-team")
MINIO_SECRET_KEY = os.getenv("MINIO_SECRET_KEY", "minio-datalabteam123")

BUCKET_GOLD    = os.getenv("BUCKET_GOLD",    "gold")
BUCKET_STAGING = os.getenv("BUCKET_STAGING", "staging")
PREFIX_PANEL   = "panel_admin"
KEY_OUTPUT     = "panel_admin/exports_gold/coherence_grades.xlsx"

s3 = boto3.client(
    "s3",
    endpoint_url          = MINIO_ENDPOINT,
    aws_access_key_id     = MINIO_ACCESS_KEY,
    aws_secret_access_key = MINIO_SECRET_KEY,
    config = Config(signature_version="s3v4", connect_timeout=30, read_timeout=120),
    region_name = "us-east-1",
    verify      = False,
)

print("\n" + "=" * 70)
print("ANALYSE DE LA COHÉRENCE DES GRADES")
print("=" * 70 + "\n")

# ============================================================
# HIÉRARCHIE DES GRADES
# ============================================================

ORDRE_GRADES = ["D1", "D2", "D3", "C1", "C2", "C3", "B1", "B3",
                "A3", "A4", "A5", "A6", "A7"]

RANG_GRADE = {g: i for i, g in enumerate(ORDRE_GRADES)}

# ============================================================
# CHARGEMENT DU PANEL GOLD
# ============================================================

print("1. Chargement des fichiers Gold panel...")

pages = s3.get_paginator("list_objects_v2").paginate(
    Bucket=BUCKET_GOLD, Prefix=PREFIX_PANEL
)
cles = sorted([
    obj["Key"]
    for page in pages
    for obj in page.get("Contents", [])
    if re.search(r"panel_20\d{2}\.parquet$", obj["Key"])
])
print(f"   Fichiers trouvés : {len(cles)}")

frames = []
for cle in cles:
    buf = io.BytesIO()
    s3.download_fileobj(BUCKET_GOLD, cle, buf)
    buf.seek(0)
    df = pl.read_parquet(buf)
    frames.append(df)
    print(f"   ✓ {os.path.basename(cle)} : {len(df):,} lignes")

base = pl.concat(frames, how="diagonal")
del frames
print(f"\n   Total : {len(base):,} lignes × {len(base.columns)} colonnes")

# Vérifications colonnes
for col in ["matricule", "mois_annee", "GRADE", "GRADE_SOURCE"]:
    if col not in base.columns:
        raise ValueError(f"Colonne '{col}' absente — vérifier le Gold panel")

# Conserver uniquement les grades valides
base = base.filter(pl.col("GRADE").is_in(ORDRE_GRADES))
print(f"   Après filtre grades valides : {len(base):,} lignes\n")

# ============================================================
# CALCUL DES TRANSITIONS
# ============================================================

print("2. Calcul des transitions de grade par agent...")

# Attribuer un rang numérique au grade
base = base.with_columns(
    pl.col("GRADE").replace(RANG_GRADE, default=None).cast(pl.Int32).alias("_rang")
)

# Trier par agent et période
base = base.sort(["matricule", "mois_annee"])

# Grade et source précédents (shift au sein de chaque agent)
base = base.with_columns([
    pl.col("GRADE").shift(1).over("matricule").alias("grade_de"),
    pl.col("GRADE_SOURCE").shift(1).over("matricule").alias("source_de"),
    pl.col("_rang").shift(1).over("matricule").alias("_rang_de"),
    pl.col("mois_annee").shift(1).over("matricule").alias("periode_de"),
])

# Garder uniquement les transitions (exclure première obs de chaque agent)
trans = base.filter(pl.col("grade_de").is_not_null()).rename({
    "GRADE":        "grade_a",
    "GRADE_SOURCE": "source_a",
    "mois_annee":   "periode_a",
})

# Amplitude du saut
trans = trans.with_columns(
    (pl.col("_rang") - pl.col("_rang_de")).alias("amplitude")
)

# Type de transition
trans = trans.with_columns(
    pl.when(pl.col("amplitude") == 0) .then(pl.lit("Stable"))
    .when(pl.col("amplitude") == 1)  .then(pl.lit("Progression normale"))
    .when(pl.col("amplitude") >  1)  .then(pl.lit("Saut en avant"))
    .otherwise(pl.lit("Régression"))
    .alias("type_transition")
)

n_trans = len(trans)
print(f"   Transitions calculées : {n_trans:,}\n")

# Distribution globale
print("   Distribution des types de transition :")
dist = trans["type_transition"].value_counts().sort("type_transition")
for row in dist.iter_rows(named=True):
    pct = 100 * row["count"] / n_trans
    print(f"     {row['type_transition']:<25} : {row['count']:>10,}  ({pct:.1f}%)")

# ============================================================
# FONCTIONS UTILITAIRES EXCEL
# ============================================================

FILL_TITRE  = PatternFill("solid", fgColor="1F4E78")
FILL_HEADER = PatternFill("solid", fgColor="4472C4")
FONT_TITRE  = Font(color="FFFFFF", bold=True, size=13)
FONT_HEADER = Font(color="FFFFFF", bold=True, size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center")


def ecrire_feuille(wb: Workbook, nom: str, df: pl.DataFrame, titre: str = "") -> None:
    ws = wb.create_sheet(nom)
    start = 1
    if titre:
        ws.merge_cells(f"A1:{get_column_letter(len(df.columns))}1")
        c = ws.cell(row=1, column=1, value=titre)
        c.fill = FILL_TITRE
        c.font = FONT_TITRE
        c.alignment = ALIGN_CENTER
        ws.row_dimensions[1].height = 22
        start = 2
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start, column=j, value=col)
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = ALIGN_CENTER
    ws.row_dimensions[start].height = 18
    for i, row in enumerate(df.iter_rows(named=True), start + 1):
        for j, col in enumerate(df.columns, 1):
            ws.cell(row=i, column=j, value=row[col])
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(len(str(col)) + 4, 35)
    ws.freeze_panes = f"A{start + 1}"
    print(f"  ✓ {nom} ({len(df):,} lignes)")


# ============================================================
# PRODUCTION DES TABLEAUX
# ============================================================

print("\n3. Production des tableaux...")

ORDRE_TYPES = ["Stable", "Progression normale", "Saut en avant", "Régression"]

wb = Workbook()
wb.remove(wb.active)

# ── 1. Résumé global : type × source ────────────────────────
# Source utilisée = source du grade d'arrivée (source_a)
resume = (
    trans
    .group_by(["type_transition", "source_a"])
    .agg(pl.len().alias("Effectif"))
    .with_columns(
        (pl.col("Effectif") * 100 / n_trans).round(2).alias("Proportion totale (%)")
    )
)

# Ajouter proportion par type (% au sein de chaque type)
totaux_type = (
    trans.group_by("type_transition")
    .agg(pl.len().alias("total_type"))
)
resume = resume.join(totaux_type, on="type_transition", how="left")
resume = resume.with_columns(
    (pl.col("Effectif") * 100 / pl.col("total_type")).round(2).alias("Proportion dans le type (%)")
).drop("total_type")

# Trier selon ordre défini
ordre_map = {t: i for i, t in enumerate(ORDRE_TYPES)}
resume = resume.with_columns(
    pl.col("type_transition").replace(ordre_map, default=99).cast(pl.Int32).alias("_ordre")
).sort(["_ordre", "source_a"]).drop("_ordre")

ecrire_feuille(wb, "RESUME_GLOBAL", resume,
               titre="Cohérence grades — Résumé par type de transition × source")

# ── 2. Transitions détaillées : grade_de × grade_à × source ─
detail = (
    trans
    .group_by(["type_transition", "grade_de", "grade_a", "source_de", "source_a"])
    .agg(pl.len().alias("Effectif"))
    .with_columns(
        (pl.col("Effectif") * 100 / n_trans).round(2).alias("Proportion totale (%)")
    )
    .sort("Effectif", descending=True)
)
ecrire_feuille(wb, "TRANSITIONS_DETAIL", detail,
               titre="Détail des transitions grade_de → grade_à × source")

# ── 3. Distribution des amplitudes de saut ──────────────────
# Uniquement sauts (amplitude > 1) et régressions (amplitude < 0)
anomalies = trans.filter(
    (pl.col("amplitude") > 1) | (pl.col("amplitude") < 0)
)
if len(anomalies) > 0:
    amplitude_stats = (
        anomalies
        .group_by(["type_transition", "amplitude", "source_a"])
        .agg(pl.len().alias("Effectif"))
        .with_columns(
            (pl.col("Effectif") * 100 / len(anomalies)).round(2).alias("Proportion anomalies (%)")
        )
        .sort(["type_transition", "amplitude"])
    )
    ecrire_feuille(wb, "AMPLITUDE_SAUTS", amplitude_stats,
                   titre="Distribution des amplitudes de saut et régressions × source")

# ============================================================
# EXPORT
# ============================================================

buf_out = io.BytesIO()
wb.save(buf_out)
buf_out.seek(0)

s3.put_object(
    Bucket=BUCKET_STAGING,
    Key=KEY_OUTPUT,
    Body=buf_out.getvalue(),
    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)

taille_kb = len(buf_out.getvalue()) / 1024
print(f"\n✅ Fichier créé : s3://{BUCKET_STAGING}/{KEY_OUTPUT}")
print(f"   Taille   : {taille_kb:.1f} KB")
print(f"   Feuilles : {len(wb.sheetnames)}\n")
print("=" * 70)
print("✓ ANALYSE COHÉRENCE GRADES TERMINÉE")
print("=" * 70)
