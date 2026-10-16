# ============================================================
# PANEL ADMIN — PLUS_AGE : LISTING DES SALAIRES AU-DESSUS D'UN SEUIL
#                              D'ANNÉES D'EXPÉRIENCE
# ============================================================
# Produit un listing détaillé (agent-mois) des salaires du périmètre de
# l'étape 11, restreint aux agents ayant PLUS de X années d'expérience.
# Le seuil X est passé en argument, ex. :  plus_age_experience.py --45
#
#   annees_experience = (période d'observation − prise_service) / 365,25
#
# Deux montants par observation (cf. 11_analyse_salaire_experience.py) :
#   M1 = montant_brut_sans_ind_sans_alloc (codes 124+228+271+425,
#        hors allocation familiale 421)
#   M2 = montant_brut_total − postes F2 461+462+463+464 (représentation)
#
# Mêmes exclusions que l'étape 11 : barèmes GREFFE, MAGISTRAT,
# para militaire, CORPS DIPL et CORPS PREF ; organisme 59 exclu.
# Le statut de conformité (Statut / Statut_Total) est conservé en clair
# dans le listing mais n'est PAS filtrant (listing exhaustif).
#
# Sources : panel_trajectoire_clean.parquet (prise_service corrigée) et
#           base_avec_statut_corrigee.parquet (salaires et statuts après 07a).
#
# Usage :
#   python plus_age_experience.py --45                 # expérience STRICTEMENT > 45
#   python plus_age_experience.py --45 --inclus        # expérience >= 45
#   python plus_age_experience.py 45                   # seuil en argument positionnel
#   python plus_age_experience.py --exp 45 --csv       # + export CSV
#   python plus_age_experience.py --45 --upload        # + publication MinIO (staging)
#
# Sortie : analyse_experience_out/salaires_plus_45_ans.xlsx
#
# Dépendances :
#   pip install polars boto3 python-dotenv openpyxl pyarrow
# ============================================================

import argparse
import os
import re
import sys
import tempfile
import time
from pathlib import Path

from dotenv import load_dotenv, find_dotenv
import boto3
from botocore.client import Config
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Cherche .env dans le dossier courant puis en remontant les parents
load_dotenv(find_dotenv(usecwd=True))

# ============================================================
# CONFIGURATION
# ============================================================
MINIO_ENDPOINT   = os.getenv("MINIO_ENDPOINT",   "http://192.168.1.230:30137")
MINIO_ACCESS_KEY = os.getenv("MINIO_ACCESS_KEY", "")
MINIO_SECRET_KEY = os.getenv("MINIO_SECRET_KEY", "")

BUCKET_GOLD    = os.getenv("BUCKET_GOLD", "gold")
BUCKET_STAGING = os.getenv("BUCKET_STAGING", "staging")

KEY_PANEL      = "panel_admin/trajectoire-pro/panel_trajectoire_clean.parquet"
KEY_STATUTS    = "panel_admin/base_avec_statut_corrigee.parquet"
PREFIX_EXPORTS = "panel_admin/exports_gold"

ANNEES          = list(range(2015, 2027))
DAYS_PER_YEAR   = 365.25
BAREMES_EXCLUS  = ("GREFFE", "MAGISTRAT", "PARA MILITAIRE", "CORPS DIPL", "CORPS PREF")

OUTDIR = Path("analyse_experience_out")   # copie locale (inspection)
CACHE  = OUTDIR / "_cache_listing_salaires.parquet"

# Colonnes du listing : (nom technique, entête lisible)
COLONNES = [
    ("matricule",             "Matricule"),
    ("mois_annee",            "Période (MMAAAA)"),
    ("annee",                 "Année"),
    ("annees_experience",     "Années d'expérience"),
    ("sexe",                  "Sexe"),
    ("situation_matrimoniale", "Situation matrimoniale"),
    ("GRADE",                 "Grade"),
    ("Code_CITP",             "Code CITP"),
    ("Metier_CITP",           "Métier (CITP)"),
    ("Statut",                "Statut M1"),
    ("m1",                    "M1 — brut sans indemnités, sans allocation"),
    ("Statut_Total",          "Statut M2"),
    ("m2",                    "M2 — brut total hors postes 461-464"),
]

COLS_NUMERIQUES = {"annee", "m1", "m2"}

# ============================================================
# CLIENT S3
# ============================================================

if not all([MINIO_ENDPOINT, MINIO_ACCESS_KEY, MINIO_SECRET_KEY]):
    raise SystemExit(
        "❌ Identifiants MinIO introuvables.\n"
        "   Créez un fichier .env dans le dossier du script (voir .env.example) :\n"
        "     MINIO_ENDPOINT, MINIO_ACCESS_KEY, MINIO_SECRET_KEY\n"
        "   ou exportez ces variables d'environnement avant de relancer."
    )

s3 = boto3.client(
    "s3",
    endpoint_url          = MINIO_ENDPOINT,
    aws_access_key_id     = MINIO_ACCESS_KEY,
    aws_secret_access_key = MINIO_SECRET_KEY,
    config                = Config(
        signature_version = "s3v4",
        retries           = {"max_attempts": 5, "mode": "adaptive"},
        connect_timeout   = 30,
        read_timeout      = 300,
    ),
    region_name = "us-east-1",
    verify      = False,
)

# ============================================================
# OUTILS POLARS
# ============================================================

def parse_date_multi(col: str) -> pl.Expr:
    s = pl.col(col).cast(pl.Utf8, strict=False)
    return pl.coalesce([
        s.str.to_date("%Y-%m-%d", strict=False),
        s.str.to_date("%d/%m/%Y", strict=False),
        s.str.to_date("%d-%m-%Y", strict=False),
        s.str.to_date("%Y/%m/%d", strict=False),
        s.str.to_date("%d.%m.%Y", strict=False),
        s.str.to_datetime(strict=False).cast(pl.Date),
    ])


def telecharger_source(key: str, nom_local: str) -> Path:
    """Télécharge une sortie existante du pipeline Gold."""
    info = s3.head_object(Bucket=BUCKET_GOLD, Key=key)
    taille = info["ContentLength"]
    tmp = Path(tempfile.gettempdir()) / nom_local
    print(f"  Téléchargement de s3://{BUCKET_GOLD}/{key}  ({taille/1024**2:.1f} MB)")
    t0 = time.time()
    s3.download_file(BUCKET_GOLD, key, str(tmp))
    print(f"  ✓ Téléchargé en {time.time()-t0:.1f}s → {tmp}")
    return tmp


# ============================================================
# CONSTRUCTION DU CACHE LOCAL COMPACT (toutes observations)
# ============================================================

def construire_cache(chemin_parquet: Path, chemin_statuts: Path) -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    t0 = time.time()

    statuts = pl.scan_parquet(str(chemin_statuts))
    noms_statuts = statuts.collect_schema().names()
    colonnes_requises = {"matricule", "mois_annee", "annee", "bareme", "Statut", "Statut_Total",
                         "montant_brut_sans_ind_sans_alloc", "montant_brut_total"}
    manquantes = colonnes_requises - set(noms_statuts)
    if manquantes:
        raise ValueError(
            f"Colonnes absentes de {KEY_STATUTS} : {sorted(manquantes)}. "
            "Relancer les étapes 06, 07 et 07a."
        )

    # Colonne organisme : l'organisme 59 est exclu (cf. 07b et étape 11).
    col_org = next((c for c in ("CODE_ORGANISME", "code_organisme", "organisme")
                    if c in noms_statuts), None)
    if col_org:
        org_expr = pl.col(col_org).cast(pl.Utf8).fill_null("(NC)").alias("_organisme")
    else:
        org_expr = pl.lit("(NC)", dtype=pl.Utf8).alias("_organisme")
        print("    ⚠️  Colonne organisme absente — exclusion de 59 impossible (M2)", flush=True)

    # M2 = montant_brut_total − postes F2 461..464 (0 si colonne absente)
    if "montant_representation_461_464" in noms_statuts:
        m2_expr = (pl.col("montant_brut_total").cast(pl.Float64)
                   - pl.col("montant_representation_461_464").cast(pl.Float64)).alias("m2")
    else:
        m2_expr = pl.col("montant_brut_total").cast(pl.Float64).alias("m2")
        print("    ⚠️  montant_representation_461_464 absente — M2 = montant_brut_total", flush=True)

    statuts = statuts.select([
        pl.col("matricule").cast(pl.Utf8),
        pl.col("mois_annee").cast(pl.Utf8),
        pl.col("bareme").cast(pl.Utf8).str.strip_chars().str.to_uppercase().alias("_bareme_filtre"),
        pl.col("Statut"), pl.col("Statut_Total"),
        pl.col("montant_brut_sans_ind_sans_alloc").cast(pl.Float64).alias("m1"),
        m2_expr,
        org_expr,
    ])

    cles = ["matricule", "mois_annee"]
    lf = (
        pl.scan_parquet(str(chemin_parquet))
        .filter(pl.col("annee").is_in(ANNEES))
        .with_columns([
            pl.col("matricule").cast(pl.Utf8),
            pl.col("mois_annee").cast(pl.Utf8),
        ])
        .join(statuts, on=cles, how="inner")
        .filter(~pl.col("_bareme_filtre").is_in(BAREMES_EXCLUS).fill_null(False))
        .filter(pl.col("_organisme") != "59")
        .with_columns([
            parse_date_multi("prise_service").alias("_ps_dt"),
            pl.date(pl.col("annee").cast(pl.Int32),
                    pl.col("mois").cast(pl.Int32), 1).alias("_periode"),
        ])
        .with_columns(
            ((pl.col("_periode") - pl.col("_ps_dt")).dt.total_days() / DAYS_PER_YEAR)
            .cast(pl.Float32).alias("annees_experience")
        )
        .select([
            pl.col("matricule").cast(pl.Utf8),
            pl.col("mois_annee").cast(pl.Utf8),
            pl.col("annee").cast(pl.Int32),
            pl.col("annees_experience"),
            pl.col("sexe").cast(pl.Utf8),
            pl.col("situation_matrimoniale").cast(pl.Utf8),
            pl.col("GRADE").cast(pl.Utf8),
            pl.col("Code_CITP").cast(pl.Utf8),
            pl.col("Metier_CITP").cast(pl.Utf8),
            pl.col("Statut"),
            pl.col("m1"),
            pl.col("Statut_Total"),
            pl.col("m2"),
        ])
    )

    lf.sink_parquet(str(CACHE), compression="zstd")

    n = pl.scan_parquet(str(CACHE)).select(pl.len()).collect().item()
    if n == 0:
        CACHE.unlink()
        raise ValueError("Aucune observation retenue après jointure et exclusions")
    print(f"  ✓ Cache écrit : {CACHE}  ({n:,} lignes, {time.time()-t0:.1f}s)")


# ============================================================
# STATISTIQUES DE RÉSUMÉ
# ============================================================

def _stats(df: pl.DataFrame, col: str) -> dict | None:
    s = df.get_column(col).drop_nulls()
    if s.len() == 0:
        return None
    return {
        "n": s.len(),
        "moyenne": round(float(s.mean()), 0),
        "mediane": round(float(s.median()), 0),
        "min": round(float(s.min()), 0),
        "max": round(float(s.max()), 0),
    }


# ============================================================
# ÉCRITURE EXCEL
# ============================================================

_FILL_TITRE = PatternFill("solid", fgColor="1F4E78")
_FILL_ENTETE = PatternFill("solid", fgColor="2E75B6")
_FONT_TITRE = Font(color="FFFFFF", bold=True, size=12)
_FONT_ENTETE = Font(color="FFFFFF", bold=True)
_ALIGN_C = Alignment(horizontal="center", vertical="center")
_FORMAT_ENTIER = '#,##0'


def _ecrire_listing(ws, df, seuil, inclus, annees):
    n_cols = len(df.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n_cols, 4))
    cell = ws.cell(row=1, column=1,
                   value=f"Salaires — agents de plus de {seuil:g} ans d'expérience"
                         f"{' (seuil inclus)' if inclus else ''} — {len(df):,} observations")
    cell.font = _FONT_TITRE
    cell.fill = _FILL_TITRE
    cell.alignment = _ALIGN_C
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n_cols, 4))
    cell = ws.cell(row=2, column=1,
                   value=f"Périmètre de l'étape 11 (barèmes exclus, organisme 59 exclu). "
                         f"Années : {annees}. M1 = brut sans indemnités/allocation ; "
                         f"M2 = brut total − postes 461-464.")
    cell.font = Font(italic=True, color="595959")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    ligne_entete = 4
    for j, (col, entete) in enumerate(COLONNES, start=1):
        c = ws.cell(row=ligne_entete, column=j, value=entete)
        c.font = _FONT_ENTETE
        c.fill = _FILL_ENTETE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[ligne_entete].height = 30

    r = ligne_entete + 1
    for ligne in df.iter_rows():
        for j, v in enumerate(ligne, start=1):
            c = ws.cell(row=r, column=j, value=v)
            nom = df.columns[j - 1]
            if nom in COLS_NUMERIQUES:
                c.number_format = _FORMAT_ENTIER
                c.alignment = Alignment(horizontal="right")
            elif nom == "annees_experience":
                c.number_format = '0.0'
                c.alignment = Alignment(horizontal="right")
        r += 1

    largeurs = {
        "matricule": 14, "mois_annee": 14, "annee": 8, "annees_experience": 18,
        "sexe": 8, "situation_matrimoniale": 20, "GRADE": 30, "Code_CITP": 14,
        "Metier_CITP": 34, "Statut": 26, "m1": 22, "Statut_Total": 26, "m2": 22,
    }
    for j, (col, _) in enumerate(COLONNES, start=1):
        ws.column_dimensions[get_column_letter(j)].width = largeurs.get(col, 16)

    ws.freeze_panes = f"A{ligne_entete + 1}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2E75B6"


def _ecrire_resume(ws, df, seuil, inclus, annees):
    ws.cell(row=1, column=1, value="RÉSUMÉ").font = _FONT_TITRE
    ws.cell(row=1, column=1).fill = _FILL_TITRE
    lignes = [
        ("Seuil", f"plus de {seuil:g} ans d'expérience" + (" (inclus)" if inclus else "")),
        ("Observations", f"{len(df):,}"),
        ("Agents distincts", f"{df.get_column('matricule').n_unique():,}"),
        ("Années couvertes", annees),
        ("Exclusions", "Barèmes GREFFE, MAGISTRAT, para militaire, CORPS DIPL, CORPS PREF ; organisme 59."),
    ]
    r = 3
    for cle, val in lignes:
        ws.cell(row=r, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=r, column=2, value=val)
        r += 1

    r += 1
    for titre, col in [("M1 — brut sans indemnités, sans allocation", "m1"),
                       ("M2 — brut total hors postes 461-464", "m2")]:
        ws.cell(row=r, column=1, value=titre).font = Font(bold=True, color="1F4E78")
        r += 1
        st = _stats(df, col)
        if st is None:
            ws.cell(row=r, column=1, value="(aucune valeur)")
            r += 1
            continue
        for k, v in st.items():
            ws.cell(row=r, column=1, value=k)
            cell = ws.cell(row=r, column=2, value=v)
            cell.number_format = _FORMAT_ENTIER if k != "n" else '#,##0'
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "8EA9C1"


def exporter(df: pl.DataFrame, seuil: float, inclus: bool, annees: str,
             chemin: Path, avec_csv: bool) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _ecrire_listing(wb.create_sheet("Salaires"), df, seuil, inclus, annees)
    _ecrire_resume(wb.create_sheet("Résumé"), df, seuil, inclus, annees)
    wb.save(chemin)
    print(f"  ✓ Excel local : {chemin}")

    if avec_csv:
        chemin_csv = chemin.with_suffix(".csv")
        df.write_csv(str(chemin_csv))
        print(f"  ✓ CSV local   : {chemin_csv}")


# ============================================================
# ARGUMENTS
# ============================================================

def _normaliser_argv(argv: list[str]) -> list[str]:
    """Accepte la forme « --45 » (ex. plus_age_experience.py --45) → « --exp 45 »."""
    out: list[str] = []
    for tok in argv:
        if re.fullmatch(r"--?\d+(?:[.,]\d+)?", tok):
            out += ["--exp", tok.lstrip("-").replace(",", ".")]
        else:
            out.append(tok)
    return out


def analyser_arguments() -> argparse.Namespace:
    ap = argparse.ArgumentParser(
        prog="plus_age_experience.py",
        description="Listing des salaires (M1/M2) pour les agents au-dessus d'un "
                    "seuil d'années d'expérience.",
        epilog="Exemples :  python plus_age_experience.py --45    |    python plus_age_experience.py --45 --inclus",
    )
    ap.add_argument("seuil_pos", nargs="?", type=float,
                    help="Seuil en années d'expérience (ex. 45 ou --45)")
    ap.add_argument("--exp", "--seuil", "--ans", dest="exp", type=float,
                    help="Seuil en années d'expérience (ex. --exp 45)")
    ap.add_argument("--inclus", action="store_true",
                    help="Inclure le seuil (expérience >= X) au lieu de strictement plus ( > X)")
    ap.add_argument("--csv", action="store_true", help="Écrire aussi une copie CSV")
    ap.add_argument("--upload", action="store_true",
                    help="Publier aussi l'Excel sur MinIO (staging)")
    ap.add_argument("--force", action="store_true", help="Reconstruire le cache local")
    ap.add_argument("--garder-parquet", action="store_true",
                    help="Ne pas supprimer les parquets téléchargés")
    args = ap.parse_args(_normaliser_argv(sys.argv[1:]))

    seuil = args.exp if args.exp is not None else args.seuil_pos
    if seuil is None:
        ap.error("précisez le seuil d'expérience, ex. :  python plus_age_experience.py --45")
    if seuil < 0:
        ap.error("le seuil doit être positif")
    args.seuil = seuil
    return args


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    args = analyser_arguments()
    seuil = args.seuil
    tag = f"{seuil:g}".replace(".", "_")
    op = ">=" if args.inclus else ">"
    fichier = OUTDIR / f"salaires_plus_{tag}_ans.xlsx"

    print("=" * 70)
    print(f" PLUS_AGE — salaires des agents avec expérience {op} {seuil:g} ans")
    print(f" Sources : s3://{BUCKET_GOLD}/{KEY_PANEL}")
    print(f"           s3://{BUCKET_GOLD}/{KEY_STATUTS}")
    print(f" Sortie  : {fichier}")
    print("=" * 70)

    # Cache local (toutes observations, seuil appliqué à la lecture)
    sources = [KEY_PANEL, KEY_STATUTS]
    dates_sources = [s3.head_object(Bucket=BUCKET_GOLD, Key=key)["LastModified"].timestamp()
                     for key in sources]
    if (CACHE.exists() and not args.force
            and CACHE.stat().st_mtime >= max(dates_sources)):
        print(f"  ↺ Cache local réutilisé : {CACHE}  (--force pour reconstruire)")
    else:
        chemins = []
        try:
            chemins.append(telecharger_source(KEY_PANEL, "panel_trajectoire_clean_pa.parquet"))
            chemins.append(telecharger_source(KEY_STATUTS, "base_avec_statut_corrigee_pa.parquet"))
            construire_cache(chemins[0], chemins[1])
        finally:
            if not args.garder_parquet:
                for chemin in chemins:
                    if chemin.exists():
                        chemin.unlink()

    # Filtrage sur le seuil + tri chronologique par agent
    lf = pl.scan_parquet(str(CACHE)).filter(pl.col("annees_experience").is_not_null())
    filtre = (pl.col("annees_experience") >= seuil) if args.inclus else (pl.col("annees_experience") > seuil)
    listing = (
        lf.filter(filtre)
        .sort(["matricule", "annee", "mois_annee"])
        .collect()
    )

    if listing.height == 0:
        raise SystemExit(f"⚠️  Aucune observation avec expérience {op} {seuil:g} ans "
                         f"dans le périmètre.")

    annees = (listing.get_column("annee").min(), listing.get_column("annee").max())
    annees_txt = f"{annees[0]} – {annees[1]}"
    n_agents = listing.get_column("matricule").n_unique()

    print(f"\n  Observations : {listing.height:,}")
    print(f"  Agents       : {n_agents:,}")
    print(f"  Années       : {annees_txt}")
    for col, lib in [("m1", "M1"), ("m2", "M2")]:
        st = _stats(listing, col)
        if st:
            print(f"  {lib} (FCFA)  : moyenne {st['moyenne']:,.0f} | médiane {st['mediane']:,.0f}"
                  f" | min {st['min']:,.0f} | max {st['max']:,.0f} | n={st['n']:,}")

    exporter(listing, seuil, args.inclus, annees_txt, fichier, args.csv)

    if args.upload:
        cle = f"{PREFIX_EXPORTS}/salaires_plus_{tag}_ans.xlsx"
        s3.upload_file(
            str(fichier), BUCKET_STAGING, cle,
            ExtraArgs={"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        )
        print(f"  ✓ Uploadé : s3://{BUCKET_STAGING}/{cle}")

    print("\n  Aperçu :")
    with pl.Config(tbl_rows=10, tbl_cols=13, fmt_str_lengths=24, tbl_width_chars=200):
        print(listing.head(10))

    print("\n" + "=" * 70)
    print(f" ✅ TERMINÉ — {listing.height:,} observations, {n_agents:,} agents")
    print("=" * 70 + "\n")


if __name__ == "__main__":
    main()
