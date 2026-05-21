# ============================================================
# PANEL ADMIN — EXPORT LISTE DES AGENTS HORS BARÈME
# ============================================================
# Lit panel_gold.parquet (colonne Statut déjà calculée par
# 07_calcul_indicateur.py) et exporte la liste nominale des
# agents hors barème.
#
# Source  : s3://gold/panel_admin/panel_gold.parquet
# Sorties :
#   s3://gold/panel_admin/liste_hors_bareme.parquet
#   s3://staging/panel_admin/exports_gold/liste_hors_bareme.xlsx
# ============================================================

import io
import os
from common.s3_client import create_s3_client
from common.logging_config import setup_pipeline_logging
import polars as pl
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ============================================================
# CONFIGURATION
# ============================================================

BUCKET_GOLD    = os.getenv("BUCKET_GOLD",    "gold")
BUCKET_STAGING = os.getenv("BUCKET_STAGING", "staging")

KEY_PANEL_GOLD = "panel_admin/panel_gold.parquet"
OUTPUT_PQ      = "panel_admin/liste_hors_bareme.parquet"
OUTPUT_XLSX    = "panel_admin/exports_gold/liste_hors_bareme.xlsx"

STATUTS_HORS_BAREME = [
    "Inférieur au minimum du barème",
    "Supérieur au maximum du barème",
]

# ============================================================
# CLIENT S3 & LOGGING
# ============================================================

s3 = create_s3_client()
logger, _ = setup_pipeline_logging("export_liste_hors_bareme")

logger.info("Démarrage de l'export liste agents hors barème")

# ============================================================
# 1. CHARGEMENT
# ============================================================

logger.info(f"1. Chargement de s3://{BUCKET_GOLD}/{KEY_PANEL_GOLD}...")

buf = io.BytesIO()
s3.download_fileobj(BUCKET_GOLD, KEY_PANEL_GOLD, buf)
buf.seek(0)
panel = pl.read_parquet(buf)

logger.info(f"   ✓ {len(panel):,} lignes × {len(panel.columns)} colonnes")

if "Statut" not in panel.columns:
    logger.error("Colonne 'Statut' absente de panel_gold.parquet")
    raise RuntimeError(
        "Colonne 'Statut' absente de panel_gold.parquet. "
        "Exécutez d'abord 07_calcul_indicateur.py puis 07b_analyse_hors_bareme.py."
    )

# ============================================================
# 2. FILTRAGE
# ============================================================

logger.info("2. Filtrage des agents hors barème...")

hors_bareme = panel.filter(pl.col("Statut").is_in(STATUTS_HORS_BAREME))
n_total = len(panel)
n_hors  = len(hors_bareme)

logger.info(f"   ✓ {n_hors:,} agents hors barème / {n_total:,} total ({100*n_hors/n_total:.1f}%)")
n_inf = int((hors_bareme["Statut"] == "Inférieur au minimum du barème").sum())
n_sup = int((hors_bareme["Statut"] == "Supérieur au maximum du barème").sum())
logger.info(f"     Inférieur au minimum : {n_inf:,}")
logger.info(f"     Supérieur au maximum : {n_sup:,}")

if n_hors == 0:
    logger.warning("Aucun agent hors barème — arrêt.")
    exit(0)

# ============================================================
# 3. EXPORT PARQUET → GOLD
# ============================================================

logger.info("3. Export Parquet → Gold...")

buf_pq = io.BytesIO()
hors_bareme.write_parquet(buf_pq, compression="snappy")
buf_pq.seek(0)
s3.put_object(Bucket=BUCKET_GOLD, Key=OUTPUT_PQ, Body=buf_pq.getvalue())
logger.info(f"   ✓ s3://{BUCKET_GOLD}/{OUTPUT_PQ}  ({len(hors_bareme):,} lignes)")

# ============================================================
# 4. EXPORT EXCEL → STAGING
# ============================================================

logger.info("4. Export Excel → Staging...")

wb = Workbook()
wb.remove(wb.active)

fill_header = PatternFill("solid", fgColor="2C3E50")
fill_inf    = PatternFill("solid", fgColor="FADBD8")   # rouge clair — inférieur au min
fill_sup    = PatternFill("solid", fgColor="FDEBD0")   # orange clair — supérieur au max
font_header = Font(color="FFFFFF", bold=True, size=10)
font_normal = Font(size=10)
al_centre   = Alignment(horizontal="center", vertical="center")
al_gauche   = Alignment(horizontal="left",   vertical="center")


def ecrire_feuille(wb: Workbook, nom: str, df: pl.DataFrame, colorier: bool = False) -> None:
    if len(df) == 0:
        return
    ws = wb.create_sheet(title=nom)
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill      = fill_header
        c.font      = font_header
        c.alignment = al_centre
    for i, row in enumerate(df.iter_rows(named=True), start=2):
        fill_ligne = None
        if colorier and "Statut" in df.columns:
            statut = row.get("Statut", "")
            if "Inférieur" in str(statut):
                fill_ligne = fill_inf
            elif "Supérieur" in str(statut):
                fill_ligne = fill_sup
        for j, col in enumerate(df.columns, start=1):
            c           = ws.cell(row=i, column=j, value=row[col])
            c.font      = font_normal
            c.alignment = al_gauche
            if fill_ligne:
                c.fill = fill_ligne
    ws.freeze_panes = "A2"
    logger.info(f"   ✓ {nom} ({len(df):,} lignes)")


# Feuille 1 : liste complète
MAX_EXCEL = 50_000
df_excel = hors_bareme.head(MAX_EXCEL)
if len(hors_bareme) > MAX_EXCEL:
    logger.warning(f"   Excel tronqué à {MAX_EXCEL:,} lignes (total dans Parquet : {n_hors:,})")
ecrire_feuille(wb, "LISTE_HORS_BAREME", df_excel, colorier=True)

# Feuille 2 : récap par Statut × Grade
col_grade = next((c for c in ["GRADE", "grade"] if c in hors_bareme.columns), None)
col_sal   = next((c for c in ["montant_brut_bareme", "montant_brut_total"] if c in hors_bareme.columns), None)
if col_grade and col_sal:
    recap_grade = (
        hors_bareme
        .group_by(["Statut", col_grade])
        .agg([
            pl.len().alias("Effectif"),
            pl.col(col_sal).mean().round(0).alias("Salaire_moyen"),
            pl.col(col_sal).min().round(0).alias("Salaire_min"),
            pl.col(col_sal).max().round(0).alias("Salaire_max"),
        ])
        .sort(["Statut", col_grade])
    )
    ecrire_feuille(wb, "RECAP_Grade", recap_grade)

# Feuille 3 : récap par Statut × Organisme
col_org = next((c for c in ["organisme", "CODE_ORGANISME"] if c in hors_bareme.columns), None)
if col_org:
    recap_org = (
        hors_bareme
        .group_by(["Statut", col_org])
        .agg([pl.len().alias("Effectif")])
        .sort(["Statut", col_org])
    )
    ecrire_feuille(wb, "RECAP_Organisme", recap_org)

buf_xlsx = io.BytesIO()
wb.save(buf_xlsx)
buf_xlsx.seek(0)
s3.put_object(
    Bucket=BUCKET_STAGING,
    Key=OUTPUT_XLSX,
    Body=buf_xlsx.getvalue(),
    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
logger.info(f"   ✓ s3://{BUCKET_STAGING}/{OUTPUT_XLSX}  ({len(buf_xlsx.getvalue()) / 1024:.1f} KB)")

# ============================================================
# RÉSUMÉ
# ============================================================

logger.info("✓ EXPORT TERMINÉ")
logger.info(f"  {n_hors:,} agents hors barème ({n_inf:,} inférieurs au min | {n_sup:,} supérieurs au max)")
logger.info(f"  Parquet : s3://{BUCKET_GOLD}/{OUTPUT_PQ}")
logger.info(f"  Excel   : s3://{BUCKET_STAGING}/{OUTPUT_XLSX}")
