# ============================================================
# PANEL ADMIN — EXPORT : SALAIRES MOYENS GLOBAUX PAR SEXE
# ============================================================
# Produit un fichier Excel (2 feuilles) :
#   Feuille 1 — Moyennes_Mensuelles     : ANNEE × MOIS × SEXE
#   Feuille 2 — Moyennes_Trimestrielles : ANNEE × TRIMESTRE × SEXE
#
# Population : agents "Conformes au barème" (hors grade D2),
#              même filtrage que 07b_calcul_indicateurs.py.
# Variable   : montant_brut_sans_ind_sans_alloc (codes 124+228+271+425)
#
# Prérequis :
#   07a_imputer_hors_bareme.py → base_avec_statut_corrigee.parquet
#
# Source  : s3://gold/panel_admin/base_avec_statut_corrigee.parquet
# Sortie  : s3://staging/panel_admin/exports_gold/salaires_moyens_par_sexe.xlsx
# ============================================================

import gc
import io
import os
from datetime import date

from common.s3_client import create_s3_client
from common.logging_config import setup_pipeline_logging

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# PARAMÈTRES
# ============================================================

INCLURE_ZEROS = False  # True = inclure salaires = 0 | False = strictement > 0
# NOTE : False aligne ce script avec 06_compiler_panel.py et superset_views.sql
#        (WHERE montant_brut_sans_ind_sans_alloc > 0) pour assurer la cohérence
#        des indicateurs entre les trois sorties.
ANNEE_MIN     = 2024   # première année incluse (inclus)

BUCKET_GOLD    = os.getenv("BUCKET_GOLD",    "gold")
BUCKET_STAGING = os.getenv("BUCKET_STAGING", "staging")
PREFIX_PANEL   = "panel_admin"
PREFIX_EXPORTS = "panel_admin/exports_gold"

KEY_BASE_CORRIGEE    = f"{PREFIX_PANEL}/base_avec_statut_corrigee.parquet"
KEY_BASE_AVEC_STATUT = f"{PREFIX_PANEL}/base_avec_statut.parquet"
OUTPUT_KEY = f"{PREFIX_EXPORTS}/salaires_moyens_par_sexe.xlsx"

VALUE_COL = "montant_brut_sans_ind_sans_alloc"

MOIS_FR = {
    1: "Janvier",  2: "Février",   3: "Mars",      4: "Avril",
    5: "Mai",      6: "Juin",      7: "Juillet",   8: "Août",
    9: "Septembre",10: "Octobre",  11: "Novembre", 12: "Décembre",
}

# ============================================================

s3 = create_s3_client()
logger, _ = setup_pipeline_logging("export_salaires_sexe")

logger.info("=" * 70)
logger.info("EXPORT — SALAIRES MOYENS GLOBAUX PAR SEXE")
logger.info("=" * 70)

# ============================================================
# 1. CHARGEMENT
# ============================================================

logger.info("1. Chargement base_avec_statut_corrigee.parquet...")

try:
    buf = io.BytesIO()
    s3.download_fileobj(BUCKET_GOLD, KEY_BASE_CORRIGEE, buf)
    buf.seek(0)
    base = pl.read_parquet(buf)
    logger.info("   Source : base_avec_statut_corrigee.parquet")
except Exception:
    logger.warning("   ⚠️  Repli sur base_avec_statut.parquet", exc_info=True)
    buf = io.BytesIO()
    s3.download_fileobj(BUCKET_GOLD, KEY_BASE_AVEC_STATUT, buf)
    buf.seek(0)
    base = pl.read_parquet(buf)

logger.info(f"   {len(base):,} lignes × {len(base.columns)} colonnes")

# ============================================================
# 2. FILTRAGE — mêmes règles que 07b_calcul_indicateurs.py
# ============================================================

logger.info("2. Filtrage...")

n_avant = len(base)
base = base.filter(pl.col("Statut") == "Conforme au barème")
logger.info(f"   Conformes au barème : {len(base):,} / {n_avant:,}")

col_grade = next((c for c in ["GRADE", "GRADE_PRINCIPAL"] if c in base.columns), None)
if col_grade:
    n_av = len(base)
    base = base.filter(pl.col(col_grade) != "D2")
    n_d2 = n_av - len(base)
    if n_d2 > 0:
        logger.info(f"   Exclusion D2 : {n_d2:,} lignes retirées")

if "SEXE_STD" not in base.columns:
    raise RuntimeError(
        "Colonne SEXE_STD absente — vérifier que 07_calcul_indicateur.py a été exécuté."
    )

# ============================================================
# 3. COLONNES TEMPORELLES
# ============================================================

if "ANNEE" not in base.columns and "mois_annee" in base.columns:
    base = base.with_columns(
        pl.col("mois_annee").str.slice(2, 4).cast(pl.Int32, strict=False).alias("ANNEE")
    )

if "MOIS_NUM" not in base.columns:
    if "mois" in base.columns:
        base = base.with_columns(pl.col("mois").alias("MOIS_NUM"))
    elif "mois_annee" in base.columns:
        base = base.with_columns(
            pl.col("mois_annee").str.slice(0, 2).cast(pl.Int32, strict=False).alias("MOIS_NUM")
        )

if "MOIS" not in base.columns and "MOIS_NUM" in base.columns:
    base = base.with_columns(
        pl.col("MOIS_NUM")
        .map_elements(lambda m: MOIS_FR.get(m, str(m) if m is not None else ""), return_dtype=pl.Utf8)
        .alias("MOIS")
    )

annee_max = date.today().year
base = base.filter(
    pl.col("ANNEE").is_not_null()
    & (pl.col("ANNEE") >= ANNEE_MIN)
    & (pl.col("ANNEE") <= annee_max)
)
logger.info(f"   Filtre années : {ANNEE_MIN} → {annee_max}  ({len(base):,} lignes retenues)")

# ============================================================
# 4. AGRÉGATION INDIVIDU × PÉRIODE (cohérent avec 07b)
# ============================================================

logger.info("3. Agrégation par agent × période...")

if "cle_unique" in base.columns:
    agg_col = "cle_unique"
elif "matricule" in base.columns:
    agg_col = "matricule"
else:
    raise RuntimeError(
        "Aucune colonne identifiant trouvée (cle_unique, matricule) — "
        "impossible de dédupliquer les agents par période."
    )

base_agrege = (
    base
    .group_by(["ANNEE", "MOIS_NUM", "MOIS", agg_col, "SEXE_STD"])
    .agg(pl.col(VALUE_COL).mean().alias(VALUE_COL))  # mean() est déterministe contrairement à first()
    .filter(
        pl.col("SEXE_STD").is_in(["Homme", "Femme"])
        & pl.col(VALUE_COL).is_not_null()
    )
)

if not INCLURE_ZEROS:
    base_agrege = base_agrege.filter(pl.col(VALUE_COL) > 0)

logger.info(f"   {len(base_agrege):,} observations (agent × période)")

# Distribution par sexe
for row in (
    base_agrege.group_by("SEXE_STD").agg(pl.len().alias("n")).sort("SEXE_STD").iter_rows(named=True)
):
    logger.info(f"     {row['SEXE_STD']:<8} : {row['n']:>10,}")

# ============================================================
# 5. STATISTIQUES PAR GROUPE
# ============================================================

def stats_par_groupe(df: pl.DataFrame, group_cols: list[str]) -> pl.DataFrame:
    return (
        df.group_by(group_cols)
        .agg([
            pl.len().alias("Effectif"),
            pl.col(VALUE_COL).mean().round(0).alias("Salaire_moyen"),
            pl.col(VALUE_COL).median().round(0).alias("Salaire_median"),
            pl.col(VALUE_COL).min().round(0).alias("Salaire_min"),
            pl.col(VALUE_COL).max().round(0).alias("Salaire_max"),
            pl.col(VALUE_COL).quantile(0.25, interpolation="nearest").round(0).alias("Salaire_P25"),
            pl.col(VALUE_COL).quantile(0.75, interpolation="nearest").round(0).alias("Salaire_P75"),
        ])
        .sort(group_cols)
    )

logger.info("4. Calcul moyennes mensuelles...")
mensuel = stats_par_groupe(base_agrege, ["ANNEE", "MOIS_NUM", "MOIS", "SEXE_STD"])
logger.info(f"   {len(mensuel):,} lignes")

logger.info("5. Calcul moyennes trimestrielles...")
base_agrege = base_agrege.with_columns(
    ((pl.col("MOIS_NUM") - 1) // 3 + 1).alias("TRIMESTRE")
)
trimestriel = (
    stats_par_groupe(base_agrege, ["ANNEE", "TRIMESTRE", "SEXE_STD"])
    .with_columns(
        pl.col("TRIMESTRE")
        .map_elements(lambda t: f"T{t}", return_dtype=pl.Utf8)
        .alias("TRIMESTRE_LABEL")
    )
    .select(["ANNEE", "TRIMESTRE", "TRIMESTRE_LABEL", "SEXE_STD",
             "Effectif", "Salaire_moyen", "Salaire_median",
             "Salaire_min", "Salaire_max", "Salaire_P25", "Salaire_P75"])
)
logger.info(f"   {len(trimestriel):,} lignes")

# ============================================================
# 6. EXPORT EXCEL
# ============================================================

logger.info("6. Export Excel...")

FILL_TITRE  = PatternFill("solid", fgColor="1F4E79")
FILL_ENTETE = PatternFill("solid", fgColor="4F81BD")
FONT_TITRE  = Font(color="FFFFFF", bold=True, size=13)
FONT_ENTETE = Font(color="FFFFFF", bold=True, size=11)
AL_CENTRE   = Alignment(horizontal="center", vertical="center")
COLS_MONETAIRES = {
    "Salaire_moyen", "Salaire_median",
    "Salaire_min", "Salaire_max", "Salaire_P25", "Salaire_P75",
}


def ecrire_feuille(
    wb: Workbook,
    nom: str,
    df: pl.DataFrame,
    titre: str,
    largeurs: dict[int, int],
) -> None:
    ws = wb.create_sheet(title=nom)
    ws.sheet_view.showGridLines = False

    n_cols = len(df.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value=titre)
    c.fill = FILL_TITRE
    c.font = FONT_TITRE
    c.alignment = AL_CENTRE
    ws.row_dimensions[1].height = 26

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=2, column=j, value=col)
        c.fill = FILL_ENTETE
        c.font = FONT_ENTETE
        c.alignment = AL_CENTRE
    ws.row_dimensions[2].height = 20

    for i, row_data in enumerate(df.iter_rows(named=True), start=3):
        for j, col in enumerate(df.columns, start=1):
            val = row_data[col]
            c = ws.cell(row=i, column=j, value=val)
            c.alignment = AL_CENTRE
            if col in COLS_MONETAIRES and val is not None:
                c.number_format = "#,##0"

    ws.freeze_panes = "A3"
    for j, width in largeurs.items():
        ws.column_dimensions[get_column_letter(j)].width = width

    logger.info(f"   ✓ {nom} ({len(df):,} lignes)")


wb = Workbook()
wb.remove(wb.active)

ecrire_feuille(
    wb, "Moyennes_Mensuelles", mensuel,
    titre="Salaires moyens par sexe — Moyennes mensuelles (conformes au barème, hors D2)",
    largeurs={1: 10, 2: 10, 3: 14, 4: 12, 5: 12,
              6: 17, 7: 18, 8: 14, 9: 14, 10: 14, 11: 14},
)

ecrire_feuille(
    wb, "Moyennes_Trimestrielles", trimestriel,
    titre="Salaires moyens par sexe — Moyennes trimestrielles (conformes au barème, hors D2)",
    largeurs={1: 10, 2: 12, 3: 16, 4: 12, 5: 12,
              6: 17, 7: 18, 8: 14, 9: 14, 10: 14, 11: 14},
)

buf_out = io.BytesIO()
wb.save(buf_out)
buf_out.seek(0)
xlsx_bytes = buf_out.getvalue()

s3.put_object(
    Bucket=BUCKET_STAGING,
    Key=OUTPUT_KEY,
    Body=xlsx_bytes,
    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)

taille_kb = len(xlsx_bytes) / 1024
logger.info(f"\n✅ Fichier créé : s3://{BUCKET_STAGING}/{OUTPUT_KEY}")
logger.info(f"   Taille   : {taille_kb:.1f} KB")
logger.info(f"   Feuilles : {len(wb.sheetnames)}")

gc.collect()
logger.info("=" * 70)
logger.info("✓ EXPORT TERMINÉ")
logger.info("=" * 70)
